from datetime import datetime
from pathlib import Path
import textwrap

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

OUT_DIR = Path(__file__).resolve().parent
FIG_DIR = OUT_DIR / "figures"
FIG_DIR.mkdir(parents=True, exist_ok=True)
DOCX_PATH = OUT_DIR / "Bao_cao_do_an_tot_nghiep_SLNA_Ticketing.docx"

FONT_REG = r"C:\Windows\Fonts\arial.ttf"
FONT_BOLD = r"C:\Windows\Fonts\arialbd.ttf"
BLUE = "#003078"
YELLOW = "#edbb00"
LIGHT = "#F4F7FB"
BORDER = "#9AA9BE"
TEXT = "#1F2937"
GREEN = "#16A34A"
RED = "#DC2626"


def font(size=24, bold=False):
    try:
        return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size=size)
    except Exception:
        return ImageFont.load_default()


def wrap_draw(draw, text, xy, fnt, fill=TEXT, width=24, line_gap=4, align="center"):
    lines = []
    for paragraph in str(text).split("\n"):
        lines.extend(textwrap.wrap(paragraph, width=width) or [""])
    x, y = xy
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=fnt)
        w = bbox[2] - bbox[0]
        xx = x - w // 2 if align == "center" else x
        draw.text((xx, y), line, font=fnt, fill=fill)
        y += (bbox[3] - bbox[1]) + line_gap
    return y


def box(draw, xy, title, subtitle="", fill="white", outline=BORDER, title_color=BLUE, wwrap=24):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=2)
    cx = (x1 + x2) // 2
    y = y1 + 16
    wrap_draw(draw, title, (cx, y), font(22, True), fill=title_color, width=wwrap)
    if subtitle:
        wrap_draw(draw, subtitle, (cx, y + 38), font(17), fill=TEXT, width=wwrap)


def arrow(draw, start, end, color=BLUE, width=3):
    import math

    draw.line([start, end], fill=color, width=width)
    ex, ey = end
    sx, sy = start
    ang = math.atan2(ey - sy, ex - sx)
    length = 14
    for delta in (2.6, -2.6):
        x = ex - length * math.cos(ang + delta)
        y = ey - length * math.sin(ang + delta)
        draw.line([(ex, ey), (x, y)], fill=color, width=width)


def save_architecture():
    img = Image.new("RGB", (1600, 920), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, 1600, 80], fill=BLUE)
    d.text((40, 22), "Sơ đồ kiến trúc tổng thể hệ thống SLNA Ticketing", font=font(30, True), fill="white")
    box(d, (80, 170, 390, 330), "Khán giả", "Trình duyệt web\nĐăng ký, đăng nhập, chọn ghế, xem vé", fill=LIGHT)
    box(d, (80, 500, 390, 660), "Quản trị viên", "Quản lý trận đấu, vé giấy, đơn hàng, nhà tài trợ", fill=LIGHT)
    box(d, (560, 130, 1010, 360), "Frontend - Next.js", "App Router, TypeScript, Ant Design, Redux\nTrang chủ, booking, checkout, my tickets, admin", fill="#FFF9DB")
    box(d, (560, 470, 1010, 700), "Backend - Express.js", "REST API, JWT Middleware, Upload, QR, Audit log\nControllers: auth, match, ticket, order, admin, sponsor, chatbot", fill="#EAF2FF")
    box(d, (1190, 130, 1500, 320), "PostgreSQL", "users, matches, tickets, orders, sponsors, audit_logs", fill="#EEFDF3")
    box(d, (1190, 450, 1500, 630), "File Upload", "Logo nhà tài trợ\nẢnh trận đấu\nTài nguyên /uploads", fill="#F8F0FF")
    box(d, (1190, 700, 1500, 850), "Dịch vụ phụ trợ", "VietQR mô phỏng\nQR vé PDF\nChatbot tra cứu thông tin CLB", fill="#FFF1F2")
    arrow(d, (390, 250), (560, 250))
    arrow(d, (390, 580), (560, 580))
    arrow(d, (785, 360), (785, 470))
    arrow(d, (1010, 570), (1190, 230))
    arrow(d, (1010, 590), (1190, 540))
    arrow(d, (1010, 620), (1190, 775))
    img.save(FIG_DIR / "architecture.png")


def save_usecase():
    img = Image.new("RGB", (1600, 1050), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, 1600, 80], fill=BLUE)
    d.text((40, 22), "Sơ đồ Use Case chính", font=font(30, True), fill="white")
    for x, label in [(125, "Khán giả"), (1425, "Quản trị viên")]:
        d.ellipse([x - 35, 180, x + 35, 250], outline=BLUE, width=4)
        d.line([x, 250, x, 390], fill=BLUE, width=4)
        d.line([x - 60, 300, x + 60, 300], fill=BLUE, width=4)
        d.line([x, 390, x - 45, 470], fill=BLUE, width=4)
        d.line([x, 390, x + 45, 470], fill=BLUE, width=4)
        d.text((x - 78, 500), label, font=font(24, True), fill=BLUE)
    d.rounded_rectangle([270, 130, 1320, 930], radius=25, outline=BORDER, width=3, fill="#FAFCFF")
    cases_left = [
        ("Đăng ký/đăng nhập", 220),
        ("Cập nhật thông tin cá nhân", 330),
        ("Xác minh CCCD", 440),
        ("Xem lịch thi đấu", 550),
        ("Chọn ghế và giữ ghế 10 phút", 660),
        ("Tạo đơn hàng và nhận QR vé", 770),
    ]
    cases_right = [
        ("Quản lý trận đấu", 220),
        ("Sinh kho vé theo khán đài", 330),
        ("Quản lý đơn hàng/doanh thu", 440),
        ("Giữ, in và bán vé giấy", 550),
        ("Soát vé bằng QR", 660),
        ("Quản lý nhà tài trợ/chatbot", 770),
    ]
    for txt, y in cases_left:
        d.ellipse([360, y - 35, 760, y + 35], fill="white", outline=BLUE, width=2)
        wrap_draw(d, txt, (560, y - 12), font(19, True), fill=TEXT, width=30)
        d.line([(185, 330), (360, y)], fill=BORDER, width=2)
    for txt, y in cases_right:
        d.ellipse([840, y - 35, 1240, y + 35], fill="white", outline=BLUE, width=2)
        wrap_draw(d, txt, (1040, y - 12), font(19, True), fill=TEXT, width=30)
        d.line([(1365, 330), (1240, y)], fill=BORDER, width=2)
    img.save(FIG_DIR / "usecase.png")


def entity(d, x, y, w, title, fields, color="#EAF2FF"):
    row_h = 30
    h = 42 + row_h * len(fields)
    d.rounded_rectangle([x, y, x + w, y + h], radius=14, fill="white", outline=BLUE, width=2)
    d.rounded_rectangle([x, y, x + w, y + 42], radius=14, fill=color, outline=BLUE, width=2)
    d.text((x + 14, y + 10), title, font=font(20, True), fill=BLUE)
    yy = y + 50
    for item in fields:
        d.text((x + 14, yy), item, font=font(16), fill=TEXT)
        yy += row_h


def save_erd():
    img = Image.new("RGB", (1800, 1250), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, 1800, 80], fill=BLUE)
    d.text((40, 22), "ERD cơ sở dữ liệu SLNA Ticketing", font=font(30, True), fill="white")
    entity(d, 80, 140, 390, "users", ["PK id", "email UNIQUE", "password", "full_name, phone_number", "role, status", "cccd, cccd_status", "auth_provider, provider_id"])
    entity(d, 700, 130, 420, "matches", ["PK id", "opponent, opponent_logo", "match_date, stadium", "competition_name", "stand_prices JSONB", "status, score", "free_stands"])
    entity(d, 80, 640, 390, "orders", ["PK id", "FK user_id -> users.id", "total_amount", "status", "payment_method", "order_qr_code", "payment_qr_code"])
    entity(d, 700, 610, 470, "tickets", ["PK id", "FK match_id -> matches.id", "FK order_id -> orders.id", "seat_code UNIQUE/match", "sector, row, seat_number", "price, status", "ticket_qr_code", "held_by, held_until", "is_printed, is_scanned"])
    entity(d, 1320, 150, 400, "sponsors", ["PK id", "name", "level", "logo_url", "website_url", "sort_order", "is_active"])
    entity(d, 1320, 650, 400, "audit_logs", ["PK id", "user_id", "action", "entity_type", "entity_id", "metadata JSONB", "created_at"])
    arrow(d, (275, 430), (275, 640))
    d.text((290, 520), "1 - N", font=font(18, True), fill=BLUE)
    arrow(d, (470, 795), (700, 795))
    d.text((555, 760), "1 - N", font=font(18, True), fill=BLUE)
    arrow(d, (910, 410), (910, 610))
    d.text((925, 505), "1 - N", font=font(18, True), fill=BLUE)
    arrow(d, (470, 260), (1320, 770), color=BORDER, width=2)
    d.text((850, 520), "user thao tác", font=font(16), fill=TEXT)
    d.text((80, 1120), "Ghi chú: ticket_status = AVAILABLE, HELD, SOLD, PAPER_RESERVED, PAPER_SOLD; order_status = PENDING, PAID, CANCELLED.", font=font(20), fill=TEXT)
    img.save(FIG_DIR / "erd.png")


def save_flow(name, title, steps, note, note_color):
    img = Image.new("RGB", (1600, 780), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, 1600, 80], fill=BLUE)
    d.text((40, 22), title, font=font(30, True), fill="white")
    x = 90
    prev = None
    for step_title, desc in steps:
        box(d, (x, 220, x + 255, 450), step_title, desc, fill="#FFF9DB", wwrap=18)
        if prev:
            arrow(d, (prev, 335), (x, 335))
        prev = x + 255
        x += 300
    d.rounded_rectangle([170, 560, 1430, 670], radius=18, fill="#FFF1F2" if note_color == RED else "#EEFDF3", outline=note_color, width=2)
    wrap_draw(d, note, (800, 595), font(24, True), fill=note_color, width=88)
    img.save(FIG_DIR / name)


def make_figures():
    save_architecture()
    save_usecase()
    save_erd()
    save_flow(
        "hold_flow.png",
        "Luồng giữ ghế và tạo đơn vé online",
        [
            ("1. Chọn ghế", "Khán giả chọn tối đa 4 ghế"),
            ("2. Giữ ghế", "AVAILABLE -> HELD\nheld_until = NOW + 10 phút"),
            ("3. Checkout", "Kiểm tra CCCD\nGiới hạn 4 vé/trận"),
            ("4. Tạo đơn", "Order PENDING\nSinh QR đơn hàng"),
            ("5. Xuất vé", "HELD -> SOLD\nSinh ticket QR"),
        ],
        "Nếu quá 10 phút chưa thanh toán, ghế HELD được đưa về AVAILABLE để người khác có thể chọn.",
        RED,
    )
    save_flow(
        "paper_flow.png",
        "Luồng quản lý vé giấy và in PDF",
        [
            ("Giữ vé giấy", "Admin chọn khán đài/số lượng\nAVAILABLE -> PAPER_RESERVED"),
            ("In vé PDF", "Sinh QR từng vé\nis_printed = true"),
            ("Bán vé giấy", "PAPER_RESERVED -> PAPER_SOLD\nGhi nhận doanh thu"),
            ("Soát vé", "Quét QR\nPAPER_SOLD hợp lệ mới được vào sân"),
        ],
        "Vé giấy đã in PDF không được trả về kênh bán online, tránh trùng vé và sai lệch kiểm soát cổng.",
        GREEN,
    )


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill.replace("#", ""))
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Calibri"
    r.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_width(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_width(table)
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
        set_cell_shading(table.rows[0].cells[i], "F2F4F7")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val)
    return table


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(85, 85, 85)


def setup_styles(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = True
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
    header = sec.header.paragraphs[0]
    header.text = "Báo cáo đồ án tốt nghiệp - SLNA Ticketing"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Hệ thống bán vé trực tuyến cho CLB Sông Lam Nghệ An")
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = RGBColor(100, 100, 100)


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TRƯỜNG ĐẠI HỌC/CAO ĐẲNG: [Điền tên trường]\nKHOA: [Điền tên khoa]")
    r.bold = True
    r.font.size = Pt(13)
    doc.add_paragraph("\n")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BÁO CÁO ĐỒ ÁN TỐT NGHIỆP")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor.from_string("003078")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Đề tài: Xây dựng hệ thống bán vé trực tuyến cho CLB Sông Lam Nghệ An")
    r.bold = True
    r.font.size = Pt(16)
    doc.add_paragraph("\n")
    add_table(
        doc,
        ["Thông tin", "Nội dung"],
        [
            ["Sinh viên thực hiện", "[Điền họ tên sinh viên]"],
            ["Mã sinh viên", "[Điền mã sinh viên]"],
            ["Lớp", "[Điền lớp]"],
            ["Giảng viên hướng dẫn", "[Điền tên giảng viên]"],
            ["Công nghệ chính", "Next.js, TypeScript, Express.js, PostgreSQL, JWT, Ant Design"],
            ["Thời gian", datetime.now().strftime("Tháng %m/%Y")],
        ],
    )
    doc.add_paragraph("\n")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Nghệ An, năm 2026").italic = True
    doc.add_page_break()


def add_front_matter(doc):
    sections = [
        ("LỜI CAM ĐOAN", "Em xin cam đoan báo cáo đồ án tốt nghiệp này được thực hiện dựa trên quá trình phân tích, thiết kế và xây dựng hệ thống bán vé trực tuyến cho CLB Sông Lam Nghệ An. Các nội dung mô tả chức năng, cơ sở dữ liệu và kiến trúc hệ thống phản ánh đúng mã nguồn của đồ án."),
        ("LỜI CẢM ƠN", "Em xin gửi lời cảm ơn tới giảng viên hướng dẫn, thầy cô trong khoa và bạn bè đã hỗ trợ trong quá trình thực hiện đồ án. Đề tài giúp em củng cố kiến thức về phát triển ứng dụng web, thiết kế cơ sở dữ liệu, bảo mật đăng nhập, quản trị hệ thống và triển khai phần mềm thực tế."),
        ("TÓM TẮT ĐỀ TÀI", "Đồ án xây dựng hệ thống bán vé trực tuyến phục vụ khán giả theo dõi các trận đấu của CLB Sông Lam Nghệ An. Hệ thống hỗ trợ đăng ký, đăng nhập, cập nhật thông tin cá nhân, xác minh CCCD, xem lịch thi đấu, chọn ghế, giữ ghế trong 10 phút, tạo đơn hàng, xuất vé QR/PDF, quản lý vé giấy, soát vé, quản lý nhà tài trợ và chatbot tra cứu thông tin."),
    ]
    for title, body in sections:
        doc.add_heading(title, level=1)
        doc.add_paragraph(body)
    doc.add_heading("MỤC LỤC", level=1)
    add_numbered(
        doc,
        [
            "Chương 1. Tổng quan đề tài",
            "Chương 2. Cơ sở lý thuyết và công nghệ sử dụng",
            "Chương 3. Phân tích và thiết kế hệ thống",
            "Chương 4. Thiết kế cơ sở dữ liệu",
            "Chương 5. Xây dựng và triển khai hệ thống",
            "Chương 6. Kiểm thử, đánh giá và hướng phát triển",
            "Kết luận và phụ lục",
        ],
    )
    doc.add_page_break()


def build_report():
    make_figures()
    doc = Document()
    setup_styles(doc)
    add_cover(doc)
    add_front_matter(doc)

    doc.add_heading("CHƯƠNG 1. TỔNG QUAN ĐỀ TÀI", level=1)
    doc.add_heading("1.1. Lý do chọn đề tài", level=2)
    doc.add_paragraph("Trong bối cảnh chuyển đổi số, việc bán vé thủ công tại sân vận động gây nhiều hạn chế như khó kiểm soát số lượng vé, dễ xảy ra trùng lặp, khó thống kê doanh thu và tốn nhân lực tại quầy. Với các trận bóng đá có lượng khán giả lớn, hệ thống bán vé trực tuyến giúp tối ưu quy trình đặt vé, giảm tải khâu phát hành, tăng tính minh bạch và hỗ trợ quản trị vận hành.")
    doc.add_paragraph("Đề tài “Xây dựng hệ thống bán vé trực tuyến cho CLB Sông Lam Nghệ An” được lựa chọn nhằm mô phỏng một bài toán thực tế: khán giả có thể xem trận đấu, chọn ghế, đặt vé và nhận QR; ban quản trị có thể quản lý lịch thi đấu, kho vé, đơn hàng, vé giấy và nhà tài trợ.")
    doc.add_heading("1.2. Mục tiêu đề tài", level=2)
    add_bullets(doc, ["Xây dựng website bán vé trực tuyến có giao diện thân thiện.", "Cho phép khán giả đăng ký, đăng nhập, cập nhật thông tin, xác minh CCCD và mua vé.", "Áp dụng cơ chế giữ ghế 10 phút để hạn chế tranh chấp khi nhiều người chọn cùng ghế.", "Tạo hệ thống quản trị cho admin quản lý trận đấu, ghế, đơn hàng, vé giấy, doanh thu và nhà tài trợ.", "Cung cấp vé QR/PDF và chức năng soát vé bằng mã QR."])
    doc.add_heading("1.3. Phạm vi đề tài", level=2)
    add_table(doc, ["Nhóm chức năng", "Nội dung thực hiện"], [["Khán giả", "Đăng ký, đăng nhập, cập nhật hồ sơ, xác minh CCCD, xem lịch, chọn ghế, giữ ghế, tạo đơn, xem vé đã mua."], ["Quản trị", "Quản lý dashboard, trận đấu, kho vé, đơn hàng, người dùng, vé giấy, soát vé, nhà tài trợ, nhật ký."], ["Tích hợp", "JWT, QR vé, upload ảnh/logo, chatbot tra cứu thông tin CLB, xuất báo cáo đơn hàng."], ["Giới hạn", "Thanh toán đang ở mức mô phỏng tạo đơn PENDING, chưa tích hợp cổng thanh toán/webhook thật."]])

    doc.add_heading("CHƯƠNG 2. CƠ SỞ LÝ THUYẾT VÀ CÔNG NGHỆ SỬ DỤNG", level=1)
    doc.add_heading("2.1. Kiến trúc ứng dụng web", level=2)
    doc.add_paragraph("Hệ thống được xây dựng theo mô hình client-server. Frontend chịu trách nhiệm giao diện và trải nghiệm người dùng, backend cung cấp REST API xử lý nghiệp vụ, còn PostgreSQL lưu trữ dữ liệu bền vững. Các request cần xác thực sử dụng JWT gửi qua header Authorization.")
    doc.add_picture(str(FIG_DIR / "architecture.png"), width=Inches(6.5))
    add_caption(doc, "Hình 2.1. Sơ đồ kiến trúc tổng thể hệ thống")
    doc.add_heading("2.2. Công nghệ sử dụng", level=2)
    add_table(doc, ["Thành phần", "Công nghệ", "Vai trò"], [["Frontend", "Next.js 16, React, TypeScript", "Xây dựng giao diện, routing, trang công khai và trang admin."], ["UI", "Ant Design, Tailwind CSS", "Tạo form, bảng, modal, layout và styling giao diện."], ["Backend", "Node.js, Express.js", "Cung cấp REST API, xử lý nghiệp vụ đặt vé, admin và chatbot."], ["Database", "PostgreSQL", "Lưu người dùng, trận đấu, vé, đơn hàng, nhà tài trợ, nhật ký."], ["Bảo mật", "JWT, bcryptjs, dotenv", "Xác thực phiên đăng nhập, mã hóa mật khẩu, quản lý secret qua .env."], ["QR/PDF", "QR code, printable ticket", "Sinh mã vé, hỗ trợ in vé online và vé giấy."]])
    doc.add_heading("2.3. Bảo mật cơ bản", level=2)
    add_bullets(doc, ["Mật khẩu người dùng được băm bằng bcrypt trước khi lưu vào database.", "JWT_SECRET, DB_PASSWORD và các thông tin nhạy cảm được đưa vào file .env.", "Middleware verifyToken kiểm tra token hợp lệ trước khi truy cập API người dùng.", "Middleware isAdmin bảo vệ các API quản trị.", "Public API tồn kho vé không trả dữ liệu doanh thu; doanh thu chỉ dành cho admin."])

    doc.add_heading("CHƯƠNG 3. PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG", level=1)
    doc.add_heading("3.1. Tác nhân hệ thống", level=2)
    add_table(doc, ["Tác nhân", "Mô tả", "Chức năng chính"], [["Khán giả", "Người dùng mua vé xem trận đấu.", "Đăng ký, đăng nhập, cập nhật thông tin, chọn ghế, mua vé, xem vé."], ["Quản trị viên", "Người vận hành hệ thống bán vé.", "Quản lý trận đấu, sinh vé, xử lý đơn, vé giấy, soát vé, nhà tài trợ, báo cáo."], ["Hệ thống", "Các tiến trình tự động và API.", "Giữ ghế 10 phút, giải phóng ghế hết hạn, sinh QR, ghi audit log."]])
    doc.add_picture(str(FIG_DIR / "usecase.png"), width=Inches(6.5))
    add_caption(doc, "Hình 3.1. Sơ đồ use case chính của hệ thống")
    doc.add_heading("3.2. Yêu cầu chức năng", level=2)
    add_table(doc, ["Mã", "Yêu cầu", "Mô tả"], [["F01", "Quản lý tài khoản", "Đăng ký, đăng nhập email/password; cập nhật thông tin; đổi mật khẩu; đổi email."], ["F02", "Xác minh CCCD", "Người dùng gửi CCCD, admin duyệt trước khi mua vé."], ["F03", "Quản lý trận đấu", "Admin tạo, sửa, xóa trận chưa bán, cập nhật tỷ số, trạng thái mở bán."], ["F04", "Chọn ghế và giữ ghế", "Khán giả chọn tối đa 4 ghế; ghế được giữ trong 10 phút."], ["F05", "Đặt vé online", "Tạo đơn PENDING, sinh QR đơn hàng và QR vé."], ["F06", "Vé giấy", "Admin giữ vé giấy, in PDF, đánh dấu bán, không trả vé đã in về online."], ["F07", "Soát vé", "Admin quét QR, kiểm tra trạng thái vé và đánh dấu đã soát."], ["F08", "Nhà tài trợ", "Admin upload logo, cập nhật URL, hiển thị logo trên trang chủ/vé PDF."], ["F09", "Chatbot", "Khán giả hỏi thông tin về CLB, trận đấu, vé và nhà tài trợ."]])
    doc.add_heading("3.3. Luồng nghiệp vụ giữ ghế và đặt vé", level=2)
    doc.add_picture(str(FIG_DIR / "hold_flow.png"), width=Inches(6.5))
    add_caption(doc, "Hình 3.2. Luồng giữ ghế và tạo đơn vé online")
    doc.add_paragraph("Khi người dùng chọn ghế, backend chuyển trạng thái ghế từ AVAILABLE sang HELD, gắn held_by với user hiện tại và held_until sau 10 phút. Nếu người dùng không hoàn tất checkout trong thời gian này, ghế được đưa về AVAILABLE.")
    doc.add_heading("3.4. Luồng nghiệp vụ vé giấy", level=2)
    doc.add_picture(str(FIG_DIR / "paper_flow.png"), width=Inches(6.5))
    add_caption(doc, "Hình 3.3. Luồng quản lý vé giấy và in PDF")

    doc.add_heading("CHƯƠNG 4. THIẾT KẾ CƠ SỞ DỮ LIỆU", level=1)
    doc.add_heading("4.1. Mô hình ERD", level=2)
    doc.add_picture(str(FIG_DIR / "erd.png"), width=Inches(6.5))
    add_caption(doc, "Hình 4.1. ERD cơ sở dữ liệu hệ thống")
    doc.add_heading("4.2. Danh sách bảng dữ liệu", level=2)
    add_table(doc, ["Bảng", "Chức năng", "Quan hệ chính"], [["users", "Lưu tài khoản người dùng/admin, thông tin cá nhân, CCCD, OAuth.", "1-N orders, liên quan audit_logs."], ["matches", "Lưu trận đấu, đội khách, thời gian, sân, giá vé từng khán đài.", "1-N tickets."], ["tickets", "Lưu từng ghế/vé, trạng thái online/vé giấy, QR, giữ ghế, in/soát vé.", "N-1 matches, N-1 orders, N-1 users."], ["orders", "Lưu đơn đặt vé, tổng tiền, trạng thái thanh toán, QR đơn hàng.", "N-1 users, 1-N tickets."], ["sponsors", "Lưu nhà tài trợ, logo, website, hạng tài trợ, trạng thái hiển thị.", "Hiển thị trên trang chủ và vé PDF."], ["audit_logs", "Ghi nhật ký thao tác quan trọng.", "Lưu user_id, action, entity_type, entity_id."]])
    doc.add_heading("4.3. Trạng thái nghiệp vụ quan trọng", level=2)
    add_table(doc, ["Đối tượng", "Trạng thái", "Ý nghĩa"], [["Ticket", "AVAILABLE", "Ghế còn trống, có thể được chọn online."], ["Ticket", "HELD", "Ghế đang được giữ bởi người dùng trong 10 phút."], ["Ticket", "SOLD", "Ghế đã được mua qua kênh online."], ["Ticket", "PAPER_RESERVED", "Ghế được admin giữ để phát hành vé giấy."], ["Ticket", "PAPER_SOLD", "Vé giấy đã bán, có thể soát bằng QR."], ["Order", "PENDING", "Đơn đã tạo, chờ xác nhận thanh toán."], ["Order", "PAID", "Đơn đã được admin xác nhận thanh toán."], ["Order", "CANCELLED", "Đơn bị hủy."]])

    doc.add_heading("CHƯƠNG 5. XÂY DỰNG VÀ TRIỂN KHAI HỆ THỐNG", level=1)
    doc.add_heading("5.1. Cấu trúc mã nguồn", level=2)
    add_table(doc, ["Thư mục/File", "Vai trò"], [["frontend/src/app", "Các page Next.js: trang chủ, lịch thi đấu, booking, checkout, profile, admin."], ["frontend/src/components", "Component dùng lại: Header, Footer, Booking, Checkout, Tickets, Admin, Chatbot."], ["frontend/src/api", "Axios client và API client cho frontend."], ["backend/routes", "Định nghĩa REST API cho auth, matches, tickets, orders, admin, sponsors, uploads, chatbot."], ["backend/controllers", "Xử lý nghiệp vụ tương ứng từng nhóm route."], ["backend/config/db.js", "Kết nối PostgreSQL qua biến môi trường."], ["backend/schema.sql", "Schema khởi tạo database."]])
    doc.add_heading("5.2. Các API chính", level=2)
    add_table(doc, ["Nhóm API", "Endpoint tiêu biểu", "Mục đích"], [["Auth", "POST /api/auth/register, POST /api/auth/login, GET /api/auth/profile", "Quản lý đăng ký, đăng nhập, hồ sơ."], ["Matches", "GET /api/matches, POST /api/matches, PUT /api/matches/:id", "Xem và quản lý trận đấu."], ["Tickets", "GET /api/tickets/:matchId, POST /api/tickets/hold/:matchId", "Lấy sơ đồ ghế, giữ/trả ghế."], ["Paper tickets", "PATCH /api/tickets/paper/:matchId, POST /api/tickets/paper/:matchId/print", "Giữ, in và quản lý vé giấy."], ["Orders", "POST /api/orders, GET /api/orders/my", "Tạo đơn và xem vé đã mua."], ["Admin", "GET /api/admin/stats, GET /api/admin/orders, GET /api/admin/users", "Dashboard, đơn hàng, người dùng, báo cáo."], ["Sponsors", "GET /api/sponsors, POST /api/sponsors", "Hiển thị và quản lý nhà tài trợ."], ["Chatbot", "POST /api/chatbot/ask", "Trả lời câu hỏi về CLB/trận đấu/vé."]])
    doc.add_heading("5.3. Module giao diện người dùng", level=2)
    add_bullets(doc, ["Trang chủ hiển thị thông tin hệ thống, trận đấu nổi bật, đếm ngược thời gian trận đấu và khu vực nhà tài trợ.", "Trang lịch thi đấu cho phép xem danh sách trận, trạng thái bán vé, giá vé và điều hướng tới trang đặt vé.", "Trang booking hiển thị sơ đồ ghế theo khán đài, trạng thái ghế trống/đã bán/đang giữ/vé giấy.", "Trang checkout hiển thị tổng tiền, phương thức thanh toán mô phỏng và kết quả tạo đơn.", "Trang vé của tôi hiển thị danh sách vé, QR và hỗ trợ in PDF theo bố cục vé điện tử.", "Trang hồ sơ cho phép cập nhật thông tin, đổi mật khẩu, đổi email và gửi CCCD."])
    doc.add_heading("5.4. Module quản trị", level=2)
    add_bullets(doc, ["Dashboard thống kê doanh thu, số vé bán, người dùng và đơn chờ xử lý.", "Quản lý trận đấu: tạo/sửa/xóa trận, upload ảnh, cấu hình giá vé từng khán đài và khán đài miễn phí.", "Quản lý kho vé: sinh ghế tự động theo khán đài A, B, C, D; xem tồn kho online/vé giấy.", "Quản lý vé giấy: giữ vé giấy, in PDF, đánh dấu bán, khóa vé đã in khỏi kênh online.", "Quản lý đơn hàng: duyệt thanh toán, hủy đơn, xuất báo cáo Excel.", "Quản lý người dùng: khóa/mở tài khoản và duyệt CCCD.", "Quản lý nhà tài trợ: upload logo, URL website, hạng tài trợ và thứ tự hiển thị.", "Soát vé: quét QR và kiểm tra trạng thái vé."])
    doc.add_heading("5.5. Cấu hình và triển khai", level=2)
    add_table(doc, ["Thành phần", "Biến/Script", "Mô tả"], [["Backend", "DB_USER, DB_HOST, DB_NAME, DB_PASSWORD, JWT_SECRET", "Bắt buộc cấu hình trong backend/.env."], ["Backend", "npm run dev", "Chạy server Express bằng nodemon."], ["Frontend", "NEXT_PUBLIC_API_BASE_URL", "Địa chỉ API backend, ví dụ http://localhost:5000/api."], ["Frontend", "npm run dev", "Chạy Next.js ở môi trường phát triển."], ["Frontend", "npm run build", "Build production bằng next build --webpack."], ["Database", "psql -f backend/schema.sql", "Khởi tạo bảng, enum và index chính."]])

    doc.add_heading("CHƯƠNG 6. KIỂM THỬ, ĐÁNH GIÁ VÀ HƯỚNG PHÁT TRIỂN", level=1)
    doc.add_heading("6.1. Kịch bản kiểm thử", level=2)
    add_table(doc, ["Mã", "Kịch bản", "Kết quả mong đợi"], [["T01", "Đăng ký tài khoản mới với email hợp lệ.", "Tạo user role USER, trạng thái ACTIVE."], ["T02", "Đăng nhập sai mật khẩu.", "API trả lỗi, không sinh access_token."], ["T03", "Người dùng chưa xác minh CCCD mua vé.", "Backend từ chối tạo đơn và yêu cầu xác minh."], ["T04", "Hai người cùng chọn một ghế.", "Chỉ một người giữ được ghế; người còn lại nhận thông báo ghế đã bị giữ."], ["T05", "Giữ ghế quá 10 phút không thanh toán.", "Ghế tự về trạng thái AVAILABLE."], ["T06", "Admin in vé giấy rồi trả về online.", "Hệ thống từ chối trả vé đã in về online."], ["T07", "Quét QR vé đã soát trước đó.", "Hệ thống cảnh báo vé đã được soát."], ["T08", "Public gọi inventory.", "Không trả dữ liệu revenue/paperRevenue."]])
    doc.add_heading("6.2. Đánh giá kết quả đạt được", level=2)
    add_bullets(doc, ["Hệ thống đáp ứng được các nghiệp vụ cốt lõi của bán vé bóng đá: trận đấu, ghế, giữ ghế, đơn hàng, vé QR, vé giấy và soát vé.", "Giao diện phân tách rõ người dùng và quản trị viên, có các trang chức năng cần thiết cho demo đồ án.", "Cơ sở dữ liệu có quan hệ rõ ràng, hỗ trợ truy vết trạng thái vé và đơn hàng.", "Các cấu hình nhạy cảm đã được đưa vào .env, hạn chế hard-code secret trong mã nguồn.", "Hệ thống có tài liệu schema và hướng dẫn chạy, thuận lợi cho hội đồng/chấm thử triển khai."])
    doc.add_heading("6.3. Hạn chế", level=2)
    add_bullets(doc, ["Thanh toán hiện mới ở mức mô phỏng tạo đơn PENDING, chưa tích hợp cổng thanh toán thật như VNPay/MoMo webhook.", "Một số migration còn được gọi trong controller để tương thích dữ liệu cũ; khi triển khai production nên tách thành hệ thống migration riêng.", "Chưa có bộ kiểm thử tự động đầy đủ cho toàn bộ backend/frontend.", "Chatbot hiện trả lời dựa trên dữ liệu/rule nội bộ, chưa tích hợp mô hình AI chuyên sâu hoặc tri thức CLB lớn."])
    doc.add_heading("6.4. Hướng phát triển", level=2)
    add_bullets(doc, ["Tích hợp cổng thanh toán thật và webhook xác nhận giao dịch tự động.", "Bổ sung email/SMS gửi vé và nhắc lịch trận đấu.", "Tối ưu sơ đồ sân theo hình dạng thực tế sân Vinh, hỗ trợ nhiều khu vực/ô cửa hơn.", "Tách migration thành công cụ riêng như Prisma, Knex hoặc file SQL versioned.", "Bổ sung unit test, integration test và kiểm thử tải cho thao tác giữ ghế đồng thời.", "Nâng cấp chatbot với dữ liệu CLB, lịch sử cầu thủ, thành tích và FAQ động."])

    doc.add_heading("KẾT LUẬN", level=1)
    doc.add_paragraph("Đồ án đã xây dựng được một hệ thống bán vé trực tuyến hoàn chỉnh ở mức demo thực tế cho CLB Sông Lam Nghệ An. Hệ thống bao gồm các chức năng dành cho khán giả, quản trị viên, quản lý vé online/vé giấy, QR/PDF, nhà tài trợ, soát vé và chatbot. Qua quá trình thực hiện, sinh viên đã vận dụng kiến thức về phân tích thiết kế hệ thống, xây dựng REST API, thiết kế cơ sở dữ liệu quan hệ, xử lý trạng thái nghiệp vụ, bảo mật JWT và phát triển giao diện web hiện đại.")
    doc.add_paragraph("Trong tương lai, hệ thống có thể tiếp tục hoàn thiện bằng cách tích hợp thanh toán thật, chuẩn hóa migration, bổ sung kiểm thử tự động và tối ưu vận hành ở môi trường production.")
    doc.add_heading("PHỤ LỤC A. DANH SÁCH TRANG FRONTEND", level=1)
    add_table(doc, ["Trang", "Vai trò"], [["/", "Trang chủ, trận đấu nổi bật, nhà tài trợ, chatbot."], ["/matches, /matches/[id]", "Danh sách và chi tiết trận đấu."], ["/booking/[id]", "Sơ đồ ghế và chọn ghế."], ["/checkout", "Thanh toán/tạo đơn."], ["/my-tickets", "Danh sách vé đã mua, in vé PDF."], ["/profile, /complete-profile", "Thông tin cá nhân, CCCD, đổi mật khẩu/email."], ["/admin/dashboard", "Thống kê doanh thu và quản trị tổng quan."], ["/admin/matches", "Quản lý trận đấu, kho vé, vé giấy."], ["/admin/orders", "Quản lý đơn hàng."], ["/admin/users", "Quản lý người dùng và duyệt CCCD."], ["/admin/sponsors", "Quản lý nhà tài trợ."], ["/admin/scanner", "Soát vé QR."]])
    doc.add_heading("PHỤ LỤC B. CÁC FILE CẤU HÌNH QUAN TRỌNG", level=1)
    add_table(doc, ["File", "Nội dung"], [["backend/.env.example", "Mẫu biến môi trường backend: database, JWT, URL frontend/backend, OAuth."], ["frontend/.env.example", "Mẫu biến môi trường frontend: NEXT_PUBLIC_API_BASE_URL."], ["backend/schema.sql", "SQL khởi tạo schema CSDL."], ["frontend/next.config.ts", "Cấu hình Next.js, image remote patterns, rewrite API fallback."], ["backend/config/db.js", "Kết nối PostgreSQL, bắt buộc đọc biến môi trường."]])

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_report()
