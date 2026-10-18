"""Structural regression checks; not a substitute for page rendering."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document

ROOT = Path(__file__).resolve().parent
PATH = ROOT / 'Bao_cao_do_an_tot_nghiep_SLNA_Ticketing_hoan_chinh.docx'

def check():
    with ZipFile(PATH) as archive:
        assert archive.testzip() is None, 'Damaged DOCX archive'
    doc = Document(PATH)
    text = '\n'.join([p.text for p in doc.paragraphs] +
                     [c.text for t in doc.tables for r in t.rows for c in r.cells])
    for stale in ('giữ ghế 15 phút', 'ghế được giữ 15 phút', 'HELD 15 phút', '8/8'):
        assert stale not in text, f'Stale claim: {stale}'
    for required in ('10 phút', '15 phút', 'SCANNER', 'SePay', 'Gemini', '49/49'):
        assert required in text, f'Missing topic: {required}'
    captions = [p.text for p in doc.paragraphs if p.style.name == 'Caption Figure']
    assert len(captions) == len(doc.inline_shapes) == 21, 'Figure/caption mismatch'
    assert len(set(captions)) == len(captions), 'Duplicate caption'
    for shape in doc.inline_shapes:
        assert shape._inline.docPr.get('descr'), 'Missing image description'
    for rel in doc.part.rels.values():
        if 'image' in rel.reltype:
            assert not rel.is_external, 'Image must be embedded'
    print(f'PASS: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, {len(captions)} figures')
    print('Layout review still requires DOCX rendering and inspection of every page.')

if __name__ == '__main__':
    check()
